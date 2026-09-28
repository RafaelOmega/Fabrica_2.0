# -*- coding: utf-8 -*-
"""Exportação do relatório de estoque geral para XLSX e CSV.

Mesmos dados e mesma ordem do PDF. XLSX grava valores numéricos com
formato de célula; CSV grava valores já formatados no padrão brasileiro
(separador ';' e BOM para o Excel).
"""
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models.relatorio_estoque_geral import LinhaEstoqueGeral
from app.reports.relatorio_estoque_geral_pdf import _moeda, _numero
from app.utils.logger import get_logger

logger = get_logger("relatorio_estoque_geral_export")

AZUL = "1F3B5B"
ZEBRA = "F2F5F8"
TOTAL = "E8EDF2"

_FMT_QTDE = "#,##0.0000"
_FMT_MOEDA = '"R$" #,##0.00'

_CABECALHOS = ["Produto", "Entradas", "Saídas", "Saldo",
               "Custo Unit.", "Valor em Estoque"]


def gerar_xlsx_estoque_geral(linhas: list[LinhaEstoqueGeral], caminho: str,
                             periodo: str = "") -> str:
    """Gera o XLSX e devolve o caminho escrito."""
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = "Estoque Geral"

    planilha.append(["RELATÓRIO DE ESTOQUE GERAL"])
    if periodo:
        planilha.append([periodo])
    planilha.append([f"Emitido em {datetime.now():%d/%m/%Y %H:%M}"])
    planilha.append([])
    planilha.append(_CABECALHOS)
    for celula in planilha[planilha.max_row]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=AZUL)

    for indice, linha in enumerate(linhas):
        planilha.append([
            f"{linha.codigo} · {linha.descricao}",
            linha.entradas,
            linha.saidas,
            linha.saldo,
            linha.custo_unitario,
            linha.valor_estoque,
        ])
        celulas = planilha[planilha.max_row]
        for celula in celulas:
            celula.number_format = _FMT_QTDE
        celulas[4].number_format = _FMT_MOEDA   # Custo Unit.
        celulas[5].number_format = _FMT_MOEDA   # Valor em Estoque
        if indice % 2:
            for celula in celulas:
                celula.fill = PatternFill("solid", fgColor=ZEBRA)

    planilha.append([
        "Total geral",
        sum(l.entradas for l in linhas),
        sum(l.saidas for l in linhas),
        None, None,
        sum(l.valor_estoque for l in linhas),
    ])
    for celula in planilha[planilha.max_row]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor=TOTAL)
        celula.number_format = _FMT_QTDE
    planilha.cell(planilha.max_row, 5).number_format = _FMT_MOEDA
    planilha.cell(planilha.max_row, 6).number_format = _FMT_MOEDA

    planilha.column_dimensions["A"].width = 45
    for coluna in ("B", "C", "D", "E", "F"):
        planilha.column_dimensions[coluna].width = 16

    pasta.save(caminho)
    logger.info("XLSX gerado: %s", caminho)
    return caminho


def gerar_csv_estoque_geral(linhas: list[LinhaEstoqueGeral], caminho: str,
                            periodo: str = "") -> str:
    """Gera o CSV e devolve o caminho escrito."""
    with open(caminho, "w", newline="", encoding="utf-8-sig") as arquivo:
        writer = csv.writer(arquivo, delimiter=";")
        writer.writerow(["RELATÓRIO DE ESTOQUE GERAL"])
        if periodo:
            writer.writerow([periodo])
        writer.writerow([f"Emitido em {datetime.now():%d/%m/%Y %H:%M}"])
        writer.writerow(_CABECALHOS)
        for linha in linhas:
            writer.writerow([
                f"{linha.codigo} · {linha.descricao}",
                _numero(linha.entradas),
                _numero(linha.saidas),
                _numero(linha.saldo),
                _moeda(linha.custo_unitario),
                _moeda(linha.valor_estoque),
            ])
        writer.writerow([
            "Total geral",
            _numero(sum(l.entradas for l in linhas)),
            _numero(sum(l.saidas for l in linhas)),
            "—", "—",
            _moeda(sum(l.valor_estoque for l in linhas)),
        ])
    logger.info("CSV gerado: %s", caminho)
    return caminho
