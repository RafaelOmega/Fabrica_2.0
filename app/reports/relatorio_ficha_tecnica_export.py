# -*- coding: utf-8 -*-
"""Exportação do relatório de fichas técnicas para XLSX e CSV.

Mesmos dados e mesma ordem do PDF. XLSX grava valores numéricos com
formato de célula; CSV grava valores já formatados no padrão brasileiro
(separador ';' e BOM para o Excel).
"""
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models.relatorio_ficha_tecnica import FichaTecnicaRelatorio
from app.reports.relatorio_ficha_tecnica_pdf import (_moeda, _moeda4, _numero,
                                                     _numero_limpo,
                                                     _valor_unitario)
from app.utils.logger import get_logger

logger = get_logger("relatorio_ficha_tecnica_export")

AZUL = "1F3B5B"
ZEBRA = "F2F5F8"
TOTAL = "E8EDF2"

_CABECALHOS = ["Código", "Insumos", "Qtde (kg)", "Custo/Saco", "Custo/kg"]
_FMT_MOEDA = '"R$" #,##0.00'
_FMT_MOEDA4 = '"R$" #,##0.0000'
_FMT_QTDE = "0.0000"


# ---------------- XLSX ----------------

def gerar_xlsx_ficha_tecnica(fichas: list[FichaTecnicaRelatorio],
                             caminho: str, periodo: str = "") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fichas Técnicas"

    ws["A1"] = "RELATÓRIO DE FICHAS TÉCNICAS"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    linha = 2
    if periodo:
        ws.cell(linha, 1, periodo)
        linha += 1
    ws.cell(linha, 1, f"Emitido em {datetime.now():%d/%m/%Y %H:%M}")
    linha += 2

    for ficha in fichas:
        linha = _bloco_xlsx(ws, linha, ficha)

    for coluna, largura in zip("ABCDE", (14, 46, 14, 14, 14)):
        ws.column_dimensions[coluna].width = largura

    wb.save(caminho)
    logger.info("XLSX gerado: %s", caminho)
    return caminho


def _bloco_xlsx(ws, linha: int,
                ficha: FichaTecnicaRelatorio) -> int:
    unitario = _valor_unitario(ficha)
    titulo = (f"Ficha {ficha.id} — {ficha.codigo_produto} · "
              f"{ficha.descricao_produto}")
    if unitario is not None:
        titulo += f" · {_moeda(unitario)}/saco"

    celula = ws.cell(linha, 1, titulo)
    celula.font = Font(bold=True, color="FFFFFF")
    celula.fill = PatternFill("solid", fgColor=AZUL)
    ws.merge_cells(start_row=linha, start_column=1,
                   end_row=linha, end_column=5)
    linha += 1

    ws.cell(linha, 1,
            f"Sacos por batida: {_numero_limpo(ficha.sacos_batida)}")
    linha += 1

    for coluna, texto in enumerate(_CABECALHOS, start=1):
        cel = ws.cell(linha, coluna, texto)
        cel.font = Font(bold=True, color=AZUL)
        cel.fill = PatternFill("solid", fgColor=ZEBRA)
    linha += 1

    if not ficha.itens:
        ws.cell(linha, 1, "—")
        ws.cell(linha, 2, "sem insumos cadastrados")
        linha += 1
    for item in ficha.itens:
        ws.cell(linha, 1, item.codigo_produto)
        ws.cell(linha, 2, item.descricao)
        ws.cell(linha, 3, item.quantidade_kg).number_format = _FMT_QTDE
        if item.custo_saco is not None:
            ws.cell(linha, 4, item.custo_saco).number_format = _FMT_MOEDA
        if item.custo_kg is not None:
            ws.cell(linha, 5, item.custo_kg).number_format = _FMT_MOEDA4
        linha += 1

    cel_rotulo = ws.cell(linha, 2, "Custo da batida")
    cel_rotulo.font = Font(bold=True)
    cel_rotulo.fill = PatternFill("solid", fgColor=TOTAL)
    cel_valor = ws.cell(linha, 5, ficha.custo_batida)
    cel_valor.font = Font(bold=True)
    cel_valor.fill = PatternFill("solid", fgColor=TOTAL)
    cel_valor.number_format = _FMT_MOEDA
    return linha + 2


# ---------------- CSV ----------------

def gerar_csv_ficha_tecnica(fichas: list[FichaTecnicaRelatorio],
                            caminho: str, periodo: str = "") -> str:
    with open(caminho, "w", newline="", encoding="utf-8-sig") as arquivo:
        writer = csv.writer(arquivo, delimiter=";")
        writer.writerow(["RELATÓRIO DE FICHAS TÉCNICAS"])
        if periodo:
            writer.writerow([periodo])
        writer.writerow([f"Emitido em {datetime.now():%d/%m/%Y %H:%M}"])
        for ficha in fichas:
            writer.writerow([])
            unitario = _valor_unitario(ficha)
            titulo = (f"Ficha {ficha.id} — {ficha.codigo_produto} · "
                      f"{ficha.descricao_produto}")
            if unitario is not None:
                titulo += f" · {_moeda(unitario)}/saco"
            writer.writerow([titulo])
            writer.writerow(
                [f"Sacos por batida: {_numero_limpo(ficha.sacos_batida)}"])
            writer.writerow(_CABECALHOS)
            if not ficha.itens:
                writer.writerow(
                    ["—", "sem insumos cadastrados", "—", "—", "—"])
            for item in ficha.itens:
                writer.writerow([
                    item.codigo_produto,
                    item.descricao,
                    _numero(item.quantidade_kg, 4),
                    _moeda(item.custo_saco)
                    if item.custo_saco is not None else "—",
                    _moeda4(item.custo_kg)
                    if item.custo_kg is not None else "—",
                ])
            writer.writerow(
                ["", "Custo da batida", "", "", _moeda(ficha.custo_batida)])
    logger.info("CSV gerado: %s", caminho)
    return caminho
