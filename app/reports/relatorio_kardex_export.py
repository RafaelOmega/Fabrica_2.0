# -*- coding: utf-8 -*-
"""Exportação do relatório de kardex para XLSX e CSV.

Mesmos dados e mesma ordem do PDF. XLSX grava valores numéricos com
formato de célula; CSV grava valores já formatados no padrão brasileiro
(separador ';' e BOM para o Excel).
"""
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models.relatorio_kardex import KardexProduto
from app.reports.relatorio_kardex_pdf import _data_iso, _numero
from app.utils.logger import get_logger

logger = get_logger("relatorio_kardex_export")

AZUL = "1F3B5B"
ZEBRA = "F2F5F8"
TOTAL = "E8EDF2"

_CABECALHOS = ["Data", "Documento", "Histórico", "Entrada", "Saída", "Saldo"]
_FMT_QTDE = "0.0000"


# ---------------- XLSX ----------------

def gerar_xlsx_kardex(kardex_list: list[KardexProduto], caminho: str,
                      periodo: str = "") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kardex"

    ws["A1"] = "RELATÓRIO DE KARDEX DO PRODUTO"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    linha = 2
    if periodo:
        ws.cell(linha, 1, periodo)
        linha += 1
    ws.cell(linha, 1, f"Emitido em {datetime.now():%d/%m/%Y %H:%M}")
    linha += 2

    for kardex in kardex_list:
        linha = _bloco_xlsx(ws, linha, kardex)

    for coluna, largura in zip("ABCDEF", (12, 12, 40, 14, 14, 14)):
        ws.column_dimensions[coluna].width = largura

    wb.save(caminho)
    logger.info("XLSX gerado: %s", caminho)
    return caminho


def _bloco_xlsx(ws, linha: int, kardex: KardexProduto) -> int:
    celula = ws.cell(linha, 1,
                     f"Produto {kardex.codigo} · {kardex.descricao}")
    celula.font = Font(bold=True, color="FFFFFF")
    celula.fill = PatternFill("solid", fgColor=AZUL)
    ws.merge_cells(start_row=linha, start_column=1,
                   end_row=linha, end_column=6)
    linha += 1

    ws.cell(linha, 1,
            f"Saldo inicial: {_numero(kardex.saldo_inicial)} kg")
    linha += 1

    for coluna, texto in enumerate(_CABECALHOS, start=1):
        cel = ws.cell(linha, coluna, texto)
        cel.font = Font(bold=True, color=AZUL)
        cel.fill = PatternFill("solid", fgColor=ZEBRA)
    linha += 1

    if not kardex.movimentos:
        ws.cell(linha, 1, "—")
        ws.cell(linha, 2, "—")
        ws.cell(linha, 3, "sem movimentos no período")
        ws.cell(linha, 6, kardex.saldo_inicial).number_format = _FMT_QTDE
        linha += 1
    for movimento in kardex.movimentos:
        ws.cell(linha, 1, _data_iso(movimento.data))
        ws.cell(linha, 2, movimento.documento)
        ws.cell(linha, 3, movimento.historico)
        if movimento.entrada:
            ws.cell(linha, 4, movimento.entrada).number_format = _FMT_QTDE
        if movimento.saida:
            ws.cell(linha, 5, movimento.saida).number_format = _FMT_QTDE
        ws.cell(linha, 6, movimento.saldo).number_format = _FMT_QTDE
        linha += 1

    ws.cell(linha, 3, "Saldo final").font = Font(bold=True)
    ws.cell(linha, 3).fill = PatternFill("solid", fgColor=TOTAL)
    for coluna, valor in ((4, kardex.entradas_total),
                          (5, kardex.saidas_total),
                          (6, kardex.saldo_final)):
        cel = ws.cell(linha, coluna, valor)
        cel.font = Font(bold=True)
        cel.fill = PatternFill("solid", fgColor=TOTAL)
        cel.number_format = _FMT_QTDE
    return linha + 2


# ---------------- CSV ----------------

def gerar_csv_kardex(kardex_list: list[KardexProduto], caminho: str,
                     periodo: str = "") -> str:
    with open(caminho, "w", newline="", encoding="utf-8-sig") as arquivo:
        writer = csv.writer(arquivo, delimiter=";")
        writer.writerow(["RELATÓRIO DE KARDEX DO PRODUTO"])
        if periodo:
            writer.writerow([periodo])
        writer.writerow([f"Emitido em {datetime.now():%d/%m/%Y %H:%M}"])
        for kardex in kardex_list:
            writer.writerow([])
            writer.writerow(
                [f"Produto {kardex.codigo} · {kardex.descricao}"])
            writer.writerow(
                [f"Saldo inicial: {_numero(kardex.saldo_inicial)} kg"])
            writer.writerow(_CABECALHOS)
            if not kardex.movimentos:
                writer.writerow(["—", "—", "sem movimentos no período",
                                 "—", "—",
                                 _numero(kardex.saldo_inicial)])
            for movimento in kardex.movimentos:
                writer.writerow([
                    _data_iso(movimento.data),
                    movimento.documento,
                    movimento.historico,
                    _numero(movimento.entrada) if movimento.entrada else "—",
                    _numero(movimento.saida) if movimento.saida else "—",
                    _numero(movimento.saldo),
                ])
            writer.writerow([
                "", "", "Saldo final",
                _numero(kardex.entradas_total),
                _numero(kardex.saidas_total),
                _numero(kardex.saldo_final),
            ])
    logger.info("CSV gerado: %s", caminho)
    return caminho
