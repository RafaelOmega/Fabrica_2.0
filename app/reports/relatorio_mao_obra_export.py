# -*- coding: utf-8 -*-
"""Exportação XLSX e CSV do relatório de mão de obra.

Padrão: detalhe + resumo. No modo somente_resumo, apenas o resumo.
"""
from datetime import datetime


def _data_br(data_iso: str) -> str:
    """AAAA-MM-DD -> DD/MM/AAAA."""
    try:
        return datetime.strptime(data_iso, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return data_iso


def gerar_xlsx_mao_obra(relatorio, caminho: str, periodo: str = "") -> None:
    """Gera o XLSX do relatório de mão de obra.

    Padrão: aba Detalhe + aba Resumo. No modo somente resumo,
    gera apenas a aba Resumo.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active

    if not relatorio.somente_resumo:
        # aba detalhe (saída a saída)
        ws.title = "Detalhe"
        ws.append(["Sequência", "Data", "Destino", "Retirada", "Código",
                   "Mão de Obra", "Qtde", "Custo", "Total"])
        for celula in ws[1]:
            celula.font = Font(bold=True)

        saida_atual = None
        total_saida = 0.0
        for linha in relatorio.linhas:
            if linha.saida_id != saida_atual:
                if saida_atual is not None:
                    ws.append(["", "", "", "", "", "Total da Saída", "", "",
                               round(total_saida, 2)])
                saida_atual = linha.saida_id
                total_saida = 0.0
            total_saida += linha.total
            ws.append([
                linha.sequencia, _data_br(linha.data_saida),
                linha.destino, linha.retirada,
                linha.codigo, linha.descricao,
                linha.quantidade, linha.custo, linha.total,
            ])
        if saida_atual is not None:
            ws.append(["", "", "", "", "", "Total da Saída", "", "",
                       round(total_saida, 2)])
        ws_resumo = wb.create_sheet("Resumo")
    else:
        # somente resumo: aba única
        ws.title = "Resumo"
        ws_resumo = ws

    # resumo por mão de obra
    ws_resumo.append(["Código", "Mão de Obra", "Qtde", "Custo", "Total"])
    for celula in ws_resumo[1]:
        celula.font = Font(bold=True)
    for linha in relatorio.resumo:
        ws_resumo.append([
            linha.codigo, linha.descricao,
            linha.quantidade, linha.custo, linha.total,
        ])
    ws_resumo.append(["", "", "", "Total geral", relatorio.total_geral])

    wb.save(caminho)


def gerar_csv_mao_obra(relatorio, caminho: str, periodo: str = "") -> None:
    """Gera o CSV do relatório de mão de obra.

    Padrão: detalhe + resumo. No modo somente resumo, apenas o resumo.
    """
    import csv

    with open(caminho, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")

        if not relatorio.somente_resumo:
            writer.writerow(["Sequência", "Data", "Destino", "Retirada",
                             "Código", "Mão de Obra", "Qtde", "Custo",
                             "Total"])
            saida_atual = None
            total_saida = 0.0
            for linha in relatorio.linhas:
                if linha.saida_id != saida_atual:
                    if saida_atual is not None:
                        writer.writerow(["", "", "", "", "",
                                         "Total da Saída", "", "",
                                         round(total_saida, 2)])
                    saida_atual = linha.saida_id
                    total_saida = 0.0
                total_saida += linha.total
                writer.writerow([
                    linha.sequencia, _data_br(linha.data_saida),
                    linha.destino, linha.retirada,
                    linha.codigo, linha.descricao,
                    linha.quantidade, linha.custo, linha.total,
                ])
            if saida_atual is not None:
                writer.writerow(["", "", "", "", "", "Total da Saída",
                                 "", "", round(total_saida, 2)])
            writer.writerow([])

        writer.writerow(["Código", "Mão de Obra", "Qtde", "Custo", "Total"])
        for linha in relatorio.resumo:
            writer.writerow([
                linha.codigo, linha.descricao,
                linha.quantidade, linha.custo, linha.total,
            ])
        writer.writerow(["", "", "", "Total geral", relatorio.total_geral])
